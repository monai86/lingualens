#!/usr/bin/env python3
"""Generate a consolidated, professionally formatted Excel workbook containing all extracted features,
model benchmark results, confounder analyses, and feature dictionaries for LinguaLens.

Output:
-------
data/ml/LinguaLens_ML_Features_and_Results_Summary.xlsx
"""

from pathlib import Path
import sys
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[2]
OUTPUT_EXCEL = PROJECT_ROOT / "data" / "ml" / "LinguaLens_ML_Features_and_Results_Summary.xlsx"

def build_excel_summary():
    print(f"Loading data files...")
    
    # 1. Load Extracted Features
    features_df = pd.read_parquet(PROJECT_ROOT / "data" / "ml" / "canonical_features_v2.parquet")
    
    # 2. Load Model Comparison Results (V1 vs V2)
    v1_v2_df = pd.read_csv(PROJECT_ROOT / "data" / "ml" / "results" / "v1_vs_v2_comparison.csv")
    
    # 3. Load Master Primary Benchmark Results
    master_df = pd.read_csv(PROJECT_ROOT / "data" / "ml" / "results" / "master_primary_results.csv")
    
    # 4. Load Age Confounding Analysis
    age_df = pd.read_csv(PROJECT_ROOT / "data" / "ml" / "results" / "age_confounding_analysis.csv")
    
    # 5. Build Feature Dictionary DataFrame
    feature_dict_data = [
        {"Category": "Metadata", "Feature Name": "participant_uid", "Type": "ID", "Description_EN": "Unique participant identifier", "Description_TH": "รหัสประจำตัวเด็ก (ไม่เปิดเผยตัวตน)"},
        {"Category": "Metadata", "Feature Name": "corpus", "Type": "Categorical", "Description_EN": "Source research corpus (Eigsti, Nadig, Flusberg, Rollins, etc.)", "Description_TH": "คลังข้อมูลวิจัยต้นทาง"},
        {"Category": "Metadata", "Feature Name": "diagnostic_group", "Type": "Categorical", "Description_EN": "Clinical category: ASD, DD (Developmental Delay), TD (Typical), SLI, etc.", "Description_TH": "กลุ่มวินิจฉัยทางคลินิก (ASD, พัฒนาการล่าช้า DD, ปกติ TD)"},
        {"Category": "Metadata", "Feature Name": "age_months", "Type": "Numeric", "Description_EN": "Age in months at time of recording", "Description_TH": "อายุของเด็กในหน่วยเดือน ณ วันที่บันทึก"},
        {"Category": "Metadata", "Feature Name": "sex", "Type": "Categorical", "Description_EN": "Biological sex (male / female)", "Description_TH": "เพศ"},
        {"Category": "Feature Schema v1", "Feature Name": "mlu", "Type": "Numeric", "Description_EN": "Mean Length of Utterance in morphemes/words (child utterances)", "Description_TH": "ความยาวประโยคเฉลี่ยของเด็ก (คำ/ประโยค) บ่งชี้ความซับซ้อนทางไวยากรณ์"},
        {"Category": "Feature Schema v1", "Feature Name": "ttr", "Type": "Numeric (0-1)", "Description_EN": "Type-Token Ratio (unique words / total words)", "Description_TH": "อัตราส่วนความหลากหลายของคำศัพท์ (คำไม่ซ้ำ / คำทั้งหมด)"},
        {"Category": "Feature Schema v1", "Feature Name": "total_words", "Type": "Integer", "Description_EN": "Total number of words spoken by the child", "Description_TH": "จำนวนคำทั้งหมดที่เด็กพูดในบทสนทนา"},
        {"Category": "Feature Schema v1", "Feature Name": "unintelligible_ratio", "Type": "Numeric (0-1)", "Description_EN": "Proportion of utterances containing unintelligible speech (xxx)", "Description_ความชัดเจน": "สัดส่วนประโยคที่ฟังไม่รู้เรื่องหรือไม่ชัดเจน"},
        {"Category": "Feature Schema v1", "Feature Name": "question_ratio", "Type": "Numeric (0-1)", "Description_EN": "Proportion of child utterances that are questions", "Description_TH": "สัดส่วนประโยคคำถามที่เด็กเป็นผู้ถาม"},
        {"Category": "Feature Schema v1", "Feature Name": "echolalia_count", "Type": "Integer", "Description_EN": "Count of immediate lexical repetitions of adult speech", "Description_TH": "จำนวนครั้งที่เด็กพูดทวนคำพูดของผู้ใหญ่ทันที"},
        {"Category": "Feature Schema v2", "Feature Name": "turn_alternation_rate", "Type": "Numeric (0-1)", "Description_EN": "Rate of adult<->child turn switches relative to total turns", "Description_TH": "อัตราการสลับรอบพูดระหว่างผู้ใหญ่กับเด็ก บ่งชี้ความลื่นไหลของการโต้ตอบ"},
        {"Category": "Feature Schema v2", "Feature Name": "speaker_balance_ratio", "Type": "Numeric (0-1)", "Description_EN": "Proportion of total session turns spoken by the child", "Description_TH": "สัดส่วนรอบพูดของเด็กเทียบกับรอบพูดทั้งหมดในเซสชัน"},
        {"Category": "Feature Schema v2", "Feature Name": "child_response_rate", "Type": "Numeric (0-1)", "Description_EN": "Proportion of adult questions/invitations answered by the child", "Description_TH": "อัตราที่เด็กตอบสนองเมื่อผู้ใหญ่ถามหรือกระตุ้น"},
        {"Category": "Feature Schema v2", "Feature Name": "adult_response_rate", "Type": "Numeric (0-1)", "Description_EN": "Proportion of child utterances responded to by the adult", "Description_TH": "อัตราที่ผู้ใหญ่ตอบสนองต่อคำพูดของเด็ก"},
        {"Category": "Feature Schema v2", "Feature Name": "partner_repetition_exact_ratio", "Type": "Numeric (0-1)", "Description_EN": "Exact turn-initial repetition of adult words (immediate echolalia)", "Description_TH": "สัดส่วนที่เด็กพูดซ้ำคำของผู้ใหญ่ตรงกันเป๊ะแบบคำต่อคำ"},
        {"Category": "Feature Schema v2", "Feature Name": "self_repetition_exact_ratio", "Type": "Numeric (0-1)", "Description_EN": "Proportion of child turns that repeat child's own immediately preceding turn", "Description_TH": "สัดส่วนที่เด็กพูดประโยคเดิมของตัวเองซ้ำๆ ติดกัน (Perseveration)"},
        {"Category": "Feature Schema v3a (Audio)", "Feature Name": "response_latency_median_ms", "Type": "Numeric (ms)", "Description_EN": "Median millisecond response latency from adult offset to child onset", "Description_TH": "ค่ามัธยฐานของเวลาตอบสนอง (มิลลิวินาที) จากผู้ใหญ่หยุดพูดจนเด็กเริ่มพูด"},
        {"Category": "Feature Schema v3a (Audio)", "Feature Name": "pitch_f0_sd_semitones", "Type": "Numeric (st)", "Description_EN": "Standard deviation of fundamental frequency F0 in semitones (ref 50Hz)", "Description_TH": "ความผันแปรของระดับเสียงพูดในหน่วย Semitones"},
        {"Category": "Feature Schema v3a (Audio)", "Feature Name": "pause_duration_ratio", "Type": "Numeric (0-1)", "Description_EN": "Proportion of within-turn pauses (>=200ms) relative to speaking duration", "Description_TH": "สัดส่วนการหยุดพัก/คิดคำในรอบพูดของเด็ก"},
        {"Category": "Audio QC", "Feature Name": "estimated_snr_db", "Type": "Numeric (dB)", "Description_EN": "Estimated continuous signal-to-noise ratio in decibels", "Description_TH": "อัตราส่วนสัญญาณเสียงต่อเสียงรบกวน (SNR ในหน่วย dB)"}
    ]
    dict_df = pd.DataFrame(feature_dict_data)
    
    print(f"Writing Excel workbook to: {OUTPUT_EXCEL}")
    
    with pd.ExcelWriter(OUTPUT_EXCEL, engine="openpyxl") as writer:
        # Sheet 1: Extracted Features
        features_df.to_excel(writer, sheet_name="Extracted_Features (1961)", index=False)
        
        # Sheet 2: Model Comparison V1 vs V2
        v1_v2_df.to_excel(writer, sheet_name="Model_Comparison_V1_vs_V2", index=False)
        
        # Sheet 3: Master Primary Benchmark
        master_df.to_excel(writer, sheet_name="Master_Primary_Benchmark", index=False)
        
        # Sheet 4: Age Confounder Analysis
        age_df.to_excel(writer, sheet_name="Age_Confounder_Analysis", index=False)
        
        # Sheet 5: Feature Dictionary
        dict_df.to_excel(writer, sheet_name="Feature_Dictionary", index=False)

    # Style the Excel Workbook with openpyxl
    wb = openpyxl.load_workbook(OUTPUT_EXCEL)
    
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    alt_fill = PatternFill(start_color="F2F7FA", end_color="F2F7FA", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    for sheet in wb.worksheets:
        # Style Header Row
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
        # Freeze top row
        sheet.freeze_panes = "A2"
        
        # Auto-adjust column widths
        for col in sheet.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col[:100]: # inspect first 100 rows for speed
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
    wb.save(OUTPUT_EXCEL)
    print(f"Successfully created polished Excel workbook at: {OUTPUT_EXCEL}")

if __name__ == "__main__":
    build_excel_summary()
